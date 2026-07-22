"""leads_import.py — Parse an uploaded leads file (CSV/XLSX) into lead rows.

Pure parsing: no DB, no network, no Redis, so it is unit-testable without any
container running. app/admin/leads.py owns the upsert side.

Column detection is deliberately NOT reimplemented here — it reuses
app/sheets/sync.py's find_column and hint words, so an uploaded file and the
Google Sheet resolve "which column holds the phone number" identically. The
codebase already had two diverging copies of that rule once, and the second one
silently failed to match any row.

Consent is never invented here. A row gets a consent basis only if the file
carries one; the endpoint may apply an operator-affirmed default afterwards.
See CLAUDE.md — a lead with no consent record is exactly as dial-ineligible as
a DND one, and due_leads() is what enforces that.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import datetime

from app.compliance.dnd import normalize_phone_e164
from app.sheets.sync import (
    CONSENT_AT_HINTS,
    CONSENT_BASIS_HINTS,
    LANGUAGE_HINTS,
    NAME_HINTS,
    PHONE_HINTS,
    find_column,
    parse_consent_at,
)

CSV_EXTENSIONS = {".csv", ".txt"}
EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}
SUPPORTED_EXTENSIONS = CSV_EXTENSIONS | EXCEL_EXTENSIONS

# A pilot-scale guard, not a real limit — it exists so a mis-picked 500k-row
# export can't tie up the event loop's thread pool or the DB for minutes.
MAX_ROWS = 5_000


class LeadsFileError(ValueError):
    """The file could not be parsed at all — wrong type, no header, no phone
    column. Distinct from a per-ROW rejection, which is reported instead."""


@dataclass
class ParsedLead:
    row_number: int
    phone_e164: str
    name: str | None = None
    language_pref: str = "auto"
    consent_basis: str | None = None
    consent_at: datetime | None = None


@dataclass
class RowError:
    row_number: int
    reason: str


@dataclass
class ParsedLeads:
    leads: list[ParsedLead] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)

    @property
    def received(self) -> int:
        return len(self.leads) + len(self.errors)


def _extension(filename: str) -> str:
    _, _, ext = filename.rpartition(".")
    return f".{ext.lower()}" if ext else ""


def _rows_from_csv(data: bytes) -> list[dict[str, object]]:
    # utf-8-sig strips the BOM Excel writes on "Save as CSV UTF-8", which
    # would otherwise become part of the first header ("﻿Phone") and stop
    # it matching any hint word.
    text = data.decode("utf-8-sig", errors="replace")
    if not text.strip():
        raise LeadsFileError("The file is empty.")
    sample = text[:8192]
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel  # single-column files give the sniffer nothing to go on
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    if not reader.fieldnames:
        raise LeadsFileError("The file has no header row.")
    return [dict(row) for row in reader]


def _rows_from_excel(data: bytes) -> list[dict[str, object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is pinned
        raise LeadsFileError(
            "Excel support requires the openpyxl package on the server."
        ) from exc

    # read_only + data_only: stream rather than build the whole object graph,
    # and take cached formula RESULTS rather than the formula text — a phone
    # column built with =CONCAT(...) would otherwise arrive as "=CONCAT(...)".
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            raise LeadsFileError("The file is empty.") from None
        headers = [str(h).strip() if h is not None else "" for h in header_row]
        if not any(headers):
            raise LeadsFileError("The file has no header row.")
        return [dict(zip(headers, r, strict=False)) for r in rows]
    finally:
        workbook.close()


def parse_leads_file(filename: str, data: bytes) -> ParsedLeads:
    """Parse *data* into leads. Raises LeadsFileError if the file as a whole is
    unusable; individual bad rows come back in .errors with their row number so
    the operator can fix them, rather than vanishing."""
    ext = _extension(filename)
    if ext in CSV_EXTENSIONS:
        rows = _rows_from_csv(data)
    elif ext in EXCEL_EXTENSIONS:
        rows = _rows_from_excel(data)
    else:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise LeadsFileError(f"Unsupported file type {ext or filename!r}. Use: {supported}.")

    if not rows:
        raise LeadsFileError("The file has a header row but no data rows.")
    if len(rows) > MAX_ROWS:
        raise LeadsFileError(
            f"The file has {len(rows)} rows; the limit is {MAX_ROWS}. Split it and upload again."
        )

    headers = [h for h in rows[0].keys() if h]
    phone_col = find_column(headers, PHONE_HINTS)
    if not phone_col:
        raise LeadsFileError(
            "No phone column found. Name a column something like "
            f"{', '.join(PHONE_HINTS)} — found: {', '.join(headers) or '(none)'}."
        )
    name_col = find_column(headers, NAME_HINTS)
    language_col = find_column(headers, LANGUAGE_HINTS)
    consent_basis_col = find_column(headers, CONSENT_BASIS_HINTS)
    consent_at_col = find_column(headers, CONSENT_AT_HINTS)

    # A "Consent At" column alone matches CONSENT_BASIS_HINTS too (it contains
    # "consent"), which would feed a timestamp in as the BASIS and leave every
    # lead permanently dial-ineligible with no warning.
    if consent_basis_col and consent_basis_col == consent_at_col:
        consent_basis_col = None

    result = ParsedLeads()
    seen: dict[str, int] = {}

    for i, row in enumerate(rows, start=2):  # row 1 is the header
        phone = normalize_phone_e164(row.get(phone_col))
        if not phone:
            raw = str(row.get(phone_col) or "").strip()
            result.errors.append(RowError(
                row_number=i,
                reason=f"not a valid Indian phone number: {raw!r}" if raw else "no phone number",
            ))
            continue
        if phone in seen:
            result.errors.append(RowError(
                row_number=i, reason=f"duplicate of row {seen[phone]} ({phone})",
            ))
            continue
        seen[phone] = i

        result.leads.append(ParsedLead(
            row_number=i,
            phone_e164=phone,
            name=(str(row.get(name_col) or "").strip() or None) if name_col else None,
            language_pref=(
                (str(row.get(language_col) or "").strip() or "auto") if language_col else "auto"
            ),
            consent_basis=(
                (str(row.get(consent_basis_col) or "").strip().lower() or None)
                if consent_basis_col else None
            ),
            consent_at=parse_consent_at(row.get(consent_at_col)) if consent_at_col else None,
        ))

    return result
